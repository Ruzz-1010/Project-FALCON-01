"""LEGACY updater for the superseded pre-Bay-Station V3 document."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCUMENT = ROOT / "THESIS DOCUMENTATION" / "PROJECT FALCON-01 - V3 Documentation.docx"


REPLACEMENTS = {
    "4. How clearly can a simplified three-page local dashboard communicate wave, environment, GPS, power, security, health, and alert information with minimal navigation for non-technical users?":
        "4. How clearly can a simplified four-page local dashboard communicate wave, environment, GPS, power, security, health, and alert information with minimal navigation for non-technical users?",
    "5. Develop a responsive dashboard with Overview, Sensors, and Logs & Alerts as the three primary pages, while keeping Buoy Motion as an optional advanced/demo visualization rather than a required operator page.":
        "5. Develop a responsive dashboard with Overview, Buoy Motion, Sensors, and Logs & Alerts as four primary navigation pages while clearly labeling Buoy Motion as an optional visualization rather than a required IMU measurement.",
    "The operator interface uses three primary pages. Detailed diagnostics remain available through expandable views rather than one card per physical sensor.":
        "The operator interface uses four primary navigation pages. Detailed diagnostics remain available through expandable views rather than one card per physical sensor.",
    "2. Optional Buoy Motion: an advanced/demo visualization driven by estimated sea context. It is not a primary navigation page and is not presented as an IMU measurement.":
        "2. Buoy Motion: an optional advanced/demo visualization driven by estimated sea context. It remains available in navigation but is not presented as an IMU measurement.",
    "2. Buoy Motion: an optional advanced/demo visualization driven by estimated sea context. It remains available in navigation but is not presented as an IMU measurement.":
        "2. Buoy Motion: an optional 3D visualization whose water-surface amplitude, heave, and tilt are generated from pressure-based estimated wave height. GPS supplies heading context only; no IMU, roll, or pitch sensor input is used.",
    "2. Buoy Motion: an optional 3D visualization whose water-surface amplitude, heave, and tilt are generated from pressure-based estimated wave height. GPS supplies heading context only; no IMU, roll, or pitch sensor input is used.":
        "2. Buoy Motion: an optional 3D visualization whose water-surface amplitude, heave, and tilt are generated from pressure-based estimated wave height. GPS supplies heading context only; no IMU, roll, or pitch sensor input is used. Current Data uses the received estimate; Calm, Moderate, Rough, and Pressure Offline are clearly labeled local presentation presets that never modify stored or live telemetry.",
    "2. Sensors: six grouped user-facing cards—Wave & Pressure, GPS & Security, Wind, Water, Power, and System. Technical details such as exact sensor model, sampling rate, calibration state, signal quality, and update age remain available through expandable details.":
        "3. Sensors: six grouped user-facing cards—Wave & Pressure, GPS & Security, Wind, Water, Power, and System. Technical details such as exact sensor model, sampling rate, calibration state, signal quality, and update age remain available through expandable details.",
    "3. Logs & Alerts: current warnings, security/calibration/operator events, persisted telemetry, acknowledgement, search, and export.":
        "4. Logs & Alerts: current warnings, security/calibration/operator events, persisted telemetry, acknowledgement, search, and export.",
    "Settings is a compact header icon. Optional predictive AI is hidden from the default operator view. The FALCON Assistant is a small animated overlay positioned near the lower edge of the interface; it may idle, wave, or move subtly and display short overlay speech bubbles. Its NORMAL, WARNING, ALERT, and OFFLINE messages explain readings, alerts, and recommended checks in simple language. It is rule-based/supporting, not a full chatbot, voice assistant, safety authority, or autonomous decision-maker.":
        "Settings is a compact header icon. Optional predictive AI is hidden from the default operator view. The FALCON Assistant design and animation files are preserved for future review, but the assistant is currently disabled and not displayed in the operator dashboard. If re-enabled, it shall remain optional, rule-based, concise, and independent of safety-critical monitoring.",
    "The repository currently includes the ESP32 firmware shell, diagnostic portal, versioned serial frame, Python simulator/edge service, SQLite storage, deterministic alerts, grouped API, four-page dashboard, optional motion visualization, optional assistant, and optional presentation prediction. The software build and automated tests demonstrate implementation behavior only. They do not prove physical sensor accuracy or coastal readiness.":
        "The repository currently includes the ESP32 firmware shell, diagnostic portal, versioned serial frame, Python simulator/edge service, SQLite storage, deterministic alerts, grouped API, four-page dashboard, optional motion visualization, saved but disabled assistant assets, and optional presentation prediction. The software build and automated tests demonstrate implementation behavior only. They do not prove physical sensor accuracy or coastal readiness.",
    "This V3.2 system-architecture-completed document supersedes conflicting V2 descriptions. The repository master context is docs/PROJECT_CONTEXT.md v6.0.":
        "This V3.3 implementation-aligned document supersedes conflicting V2 and earlier V3 descriptions. The repository master context is docs/PROJECT_CONTEXT.md v6.1. The physical and visual prototype is under redesign; old CAD geometry and placement remain historical references until adviser approval.",
    "The FALCON Assistant appears as a small animated overlay near the lower edge of the dashboard. It may perform simple idle/wave animations and show short text bubbles such as system-normal summaries, calibration reminders, sensor-offline notices, low-battery messages, or security alerts. The assistant translates existing system states into simpler language; it does not create independent safety decisions.":
        "The FALCON Assistant concept, mascot, and animation files are retained for possible future use, but the assistant is currently hidden from the dashboard while its final operator design is reviewed. If approved and re-enabled, it may translate existing validated system states into short explanations; it shall not create independent safety decisions.",
    "The default operator interface is intentionally simplified for quick use by non-technical and older users. The three primary pages are Overview, Sensors, and Logs & Alerts. Important information should be understandable with minimal clicking, large readable labels, high contrast, and clear status colors.":
        "The default operator interface is intentionally simplified for quick use by non-technical and older users. The four navigation pages are Overview, Buoy Motion, Sensors, and Logs & Alerts; Buoy Motion is explicitly optional visualization. Important information should be understandable with minimal clicking, large readable labels, high contrast, and clear status colors.",
    "Sensors → ESP32 acquisition/validation → USB serial/UART → Orange Pi local processing/storage/API → local dashboard → user. Optional FALCON Assistant consumes already validated system state and does not control or replace the core monitoring pipeline.":
        "Sensors → ESP32 acquisition/validation → USB serial/UART → Orange Pi local processing/storage/API → local dashboard → user. The saved FALCON Assistant is currently disabled; if re-enabled, it may consume validated system state but shall never control or replace the core monitoring pipeline.",
    "Provide simplified three-page browser dashboard.": "Provide simplified four-page browser dashboard.",
    "Display current state, sensors, logs, alerts, and optional assistant": "Display current state, optional motion visualization, grouped sensors, logs, and alerts",
    "optional motion visualization": "optional pressure-driven motion visualization",
    "It does not claim laboratory-grade salinity": "It does not claim laboratory-grade water-quality analysis",
    "3. Bench-integrate Bar02, GPS, wind, DS18B20, conductivity, health, and security channels.":
        "3. Confirm and bench-integrate the HPT604 4–20 mA pressure loop, GPS, wind, health, and security channels; retain Bar02 only for short supervised comparison.",
    "conductivity/salinity hardware; ": "",
}


def replace_text(text: str) -> str:
    for old, new in REPLACEMENTS.items():
        text = text.replace(old, new)
    text = text.replace("docs/PROJECT_CONTEXT.md v6.0", "docs/PROJECT_CONTEXT.md v6.1")
    if text.startswith("This V3.3 implementation-aligned document supersedes") and "under redesign" not in text:
        text += (
            " The physical and visual prototype is under redesign. Old CAD dimensions, component placement, "
            "solar arrangement, cooling geometry, and renders are historical references until adviser approval."
        )
    return text


def paragraph_start(document: Document, prefix: str):
    return next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix))


def add_table(document: Document, headings: list[str], rows: list[list[str]]):
    table = document.add_table(rows=1, cols=len(headings))
    if len(document.tables) > 1:
        table.style = document.tables[0].style
    for index, heading in enumerate(headings):
        table.rows[0].cells[index].text = heading
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    return table


def update() -> None:
    document = Document(DOCUMENT)
    for paragraph in document.paragraphs:
        revised = replace_text(paragraph.text)
        if revised != paragraph.text:
            paragraph.text = revised

    version = document.paragraphs[7]
    version.text = "V3.3 — CURRENT IMPLEMENTATION ALIGNED\nUndergraduate Thesis Documentation\n26 August 2026"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    revised = replace_text(paragraph.text)
                    if revised != paragraph.text:
                        paragraph.text = revised

    for table in document.tables:
        for row in list(table.rows):
            if any("Conductivity/salinity" in cell.text for cell in row.cells):
                table._tbl.remove(row._tr)

    wave_anchor = paragraph_start(document, "8. Pressure-Based")
    if not any(paragraph.text.startswith("7.1 Water Temperature") for paragraph in document.paragraphs):
        subsection_style = paragraph_start(document, "4.1 General Objective").style
        water_elements = []
        water_heading = document.add_paragraph("7.1 Water Temperature Monitoring")
        water_heading.style = subsection_style
        water_elements.append(water_heading._p)
        water_intro = document.add_paragraph(
            "FALCON includes a sealed DS18B20 supporting sensor that reports water temperature in degrees Celsius. "
            "This channel provides environmental context and does not constitute laboratory-grade water-quality analysis."
        )
        water_elements.append(water_intro._p)
        water_table = add_table(document, ["Channel", "Required output", "Validity and calibration rule"], [
            ["Sealed DS18B20", "Water temperature (°C), timestamp, validity, and freshness", "Compare with a traceable reference thermometer; report OFFLINE or STALE when invalid"],
        ])
        water_elements.append(water_table._tbl)
        water_note = document.add_paragraph(
            "The dashboard displays this channel under Water. Exact model identification, sampling, freshness, validity, and reference-comparison status remain available in expanded technical details."
        )
        water_elements.append(water_note._p)
        for element in water_elements:
            wave_anchor._p.addprevious(element)

    ordered = [
        paragraph_start(document, "18. Pressure Sensor Installation"),
        paragraph_start(document, "Wave-height accuracy shall"),
        paragraph_start(document, "The exact installation depth"),
        paragraph_start(document, "The pressure sensor shall"),
        paragraph_start(document, "19. Simplified Operator Interface"),
        paragraph_start(document, "The Sensors page does not"),
        paragraph_start(document, "The default operator interface"),
        paragraph_start(document, "The FALCON Assistant concept"),
        paragraph_start(document, "20. Literature Basis"),
        paragraph_start(document, "The pressure-based method is supported"),
        paragraph_start(document, "For Project FALCON, these studies"),
        paragraph_start(document, "21. Reference Literature"),
        paragraph_start(document, "Bonneton, P."),
        paragraph_start(document, "Bishop, C."),
        paragraph_start(document, "Ardhuin, F."),
    ]
    anchor = paragraph_start(document, "22. Power Architecture")
    for paragraph in ordered:
        anchor._p.addprevious(paragraph._p)

    power_anchor = paragraph_start(document, "Figure 5. Solar")
    if not any(paragraph.text.startswith("22.1 Provisional Electronics") for paragraph in document.paragraphs):
        additions = []
        subsection_style = paragraph_start(document, "4.1 General Objective").style
        section = document.add_paragraph("22.1 Provisional Electronics Consumption")
        section.style = subsection_style
        additions.append(section._p)
        intro = document.add_paragraph(
            "The following values are design envelopes and planning assumptions, not measured continuous consumption. "
            "Final values shall be replaced by a 24-hour current log covering Orange Pi startup, Wi-Fi activity, storage writes, sensor sampling, fan startup, and idle operation."
        )
        additions.append(intro._p)
        load_table = add_table(document, ["Load", "Preliminary electrical basis", "Documentation status"], [
            ["Orange Pi Zero 3, storage, and Wi-Fi", "Regulated 5 V / 3 A supply envelope (15 W maximum available, not assumed average draw)", "Measure startup peak, idle, Wi-Fi, and database-write current"],
            ["ESP32 and complete sensor carrier", "Separate regulated 5 V / 2 A branch envelope (10 W maximum available)", "Measure ESP32 plus every installed sensor; do not use the envelope as average consumption"],
            ["Two optional enclosure fans", "2 × 5 V × 0.15 A = 1.5 W maximum", "Include only if the selected fans are installed and thermally required"],
            ["Pressure, GPS, wind, water, power, and security sensors", "Included in the ESP32/sensor branch", "Record individual idle/active current from exact purchased datasheets and bench measurements"],
            ["DC conversion and charging losses", "Included through conservative system-efficiency assumptions", "Verify buck and MPPT efficiency across battery voltage, load, and enclosure temperature"],
        ])
        additions.append(load_table._tbl)
        section = document.add_paragraph("22.2 Battery Capacity and Autonomy")
        section.style = subsection_style
        additions.append(section._p)
        battery_text = document.add_paragraph(
            "The provisional baseline is a 12.8 V, 20 Ah LiFePO4 battery: 12.8 V × 20 Ah = 256 Wh nominal. "
            "Using an 80% usable-energy planning limit gives 204.8 Wh. Estimated no-solar runtime is usable energy divided by measured average load."
        )
        additions.append(battery_text._p)
        battery_table = add_table(document, ["Complete average load", "Daily energy", "Estimated no-solar runtime"], [
            ["6 W", "144 Wh/day", "34.1 hours"],
            ["8 W", "192 Wh/day", "25.6 hours"],
            ["10 W", "240 Wh/day", "20.5 hours"],
        ])
        additions.append(battery_table._tbl)
        section = document.add_paragraph("22.3 Solar-Panel Sizing")
        section.style = subsection_style
        additions.append(section._p)
        solar_text = document.add_paragraph(
            "Preliminary daily harvest uses panel rating × 4 peak-sun-hours × 70% net system efficiency. "
            "The 70% factor provisionally covers MPPT, wiring, conversion, temperature, orientation, and contamination losses; actual Puerto Princesa conditions must be logged."
        )
        additions.append(solar_text._p)
        solar_table = add_table(document, ["Solar option", "Estimated harvest", "Interpretation"], [
            ["60 W nominal panel", "168 Wh/day", "Supervised-test minimum; supports a 6 W average load with only about 24 Wh/day theoretical margin"],
            ["80 W nominal panel", "224 Wh/day", "Preferred prototype starting point; about 80 Wh/day margin at 6 W or 32 Wh/day at 8 W"],
            ["LiFePO4-compatible MPPT", "Exact model TBD", "Panel Voc/Isc, charge profile, current rating, thermal behavior, and protection must be verified"],
        ])
        additions.append(solar_table._tbl)
        warning = document.add_paragraph(
            "These calculations are provisional sizing values, not validated endurance claims. Do not approve unattended deployment until the system passes a 24-hour load measurement, converter/brownout tests, and at least a 72-hour solar-endurance trial."
        )
        additions.append(warning._p)
        for element in additions:
            power_anchor._p.addprevious(element)

    requirement_table = next((table for table in document.tables if any(cell.text == "FR-01" for row in table.rows for cell in row.cells)), None)
    if requirement_table is not None and not any(cell.text == "FR-08" for row in requirement_table.rows for cell in row.cells):
        row = requirement_table.add_row().cells
        values = [
            "FR-08",
            "Acquire sealed-probe water temperature with explicit validity and freshness states.",
            "DS18B20 reference comparison, stale/disconnect tests, and correct dashboard labels.",
            "Required",
        ]
        for index, value in enumerate(values):
            row[index].text = value

    document.core_properties.title = "Project FALCON-01 - V3.3 Documentation"
    document.core_properties.subject = "Current implementation-aligned pressure-based coastal monitoring buoy documentation"
    document.core_properties.comments = "Aligned with repository master context v6.1 and dashboard implementation on 2026-08-27. Physical prototype under redesign."
    temporary = DOCUMENT.with_suffix(".updated.docx")
    document.save(temporary)
    temporary.replace(DOCUMENT)
    print(DOCUMENT)


if __name__ == "__main__":
    raise SystemExit("Legacy updater blocked: use update_baystation_docx.py")
