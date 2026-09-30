"""Stamp every Project FALCON DOCX with the current pressure-sensor decision.

The canonical thesis is regenerated separately. This script adds a conspicuous,
idempotent document-control notice to all supporting and legacy DOCX records so
older Bar02 text cannot silently override the 2026-09-10 engineering baseline.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
THESIS = ROOT / "THESIS DOCUMENTATION"
MARKER = "PROJECT FALCON DOCUMENT-CONTROL UPDATE — 10 SEPTEMBER 2026"

LEGACY_NAMES = {
    "Falcon v3 documentation.docx",
    "FalconRRL.docx",
    "PROJECT FALCON-01 -  V2 REVISED.docx",
    "PROJECT FALCON-01 - V3 Documentation - LEGACY.docx",
    "PROJECT FALCON-01.docx",
    "Project_FALCON_Full_Documentation.docx",
    "Project_FALCON_Sensor_List.docx",
    "Project_FALCON_Sensor_Selection_and_Architecture.docx",
    "Project_FALCON_Sensor_Validation_and_Calibration_Test_Plan.docx",
    "Project_FALCON_Wednesday_Presenter_Guide.docx",
}

CANONICAL_NAMES = {
    "BayStation.docx",
    "PROJECT FALCON-01 - V3 Documentation.docx",
}

DECISION = (
    "Current pressure-sensor baseline: Holykell HPT604 Type A, provisionally "
    "0–2 mH2O vented gauge with 4–20 mA output, is the recommended long-duration "
    "deployment candidate. The exact order code, continuous-saltwater suitability, "
    "wetted materials, seal, cable and calibration must be confirmed before purchase. "
    "The protected 12 V loop requires a 150 ohm 0.1% shunt, input protection/filtering "
    "and a 3.3 V ADS1115 receiver. Bar02 is bench-only; its old I2C PCB and wiring are "
    "not fabrication-ready. No hardware is represented as purchased or field-validated."
)


def remove_old_notice(document: Document) -> None:
    for paragraph in list(document.paragraphs):
        if paragraph.text.startswith(MARKER) or paragraph.text.startswith("STATUS OF THIS COPY:") or paragraph.text.startswith("Current pressure-sensor baseline:"):
            paragraph._element.getparent().remove(paragraph._element)


def prepend_paragraph(document: Document, text: str, *, bold=False, color="1F2937", size=9, shade=None) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if shade:
        properties = paragraph._p.get_or_add_pPr()
        fill = OxmlElement("w:shd")
        fill.set(qn("w:fill"), shade)
        properties.append(fill)
    body = document._element.body
    body.remove(paragraph._p)
    body.insert(0, paragraph._p)


def update(path: Path) -> None:
    document = Document(path)
    remove_old_notice(document)
    if path.name in CANONICAL_NAMES:
        status = "STATUS OF THIS COPY: CURRENT CANONICAL THESIS — Version 4.1, generated from the Project FALCON v8.2 baseline."
    elif path.name in LEGACY_NAMES:
        status = (
            "STATUS OF THIS COPY: LEGACY / SUPERSEDED REFERENCE. Use BayStation.docx and "
            "docs/PROJECT_CONTEXT.md for the current thesis and architecture."
        )
    else:
        status = "STATUS OF THIS COPY: CURRENT SUPPORTING RECORD, subject to the canonical BayStation.docx and docs/PROJECT_CONTEXT.md."
    # Insert in reverse order because every new paragraph is moved to body index 0.
    prepend_paragraph(document, DECISION, color="374151", size=8, shade="F3F4F6")
    prepend_paragraph(document, status, bold=True, color="B45309", size=9, shade="FEF3C7")
    prepend_paragraph(document, MARKER, bold=True, color="FFFFFF", size=10, shade="0F766E")
    document.core_properties.comments = (
        "Pressure-sensor baseline synchronized 2026-09-10. HPT604 is a candidate; "
        "Bar02 is bench-only. Planned hardware remains unvalidated."
    )
    document.save(path)
    print(path.relative_to(ROOT))


def main() -> None:
    paths = sorted(THESIS.rglob("*.docx"))
    if not paths:
        raise SystemExit("No DOCX files found")
    for path in paths:
        update(path)
    print(f"Updated {len(paths)} DOCX files")


if __name__ == "__main__":
    main()
