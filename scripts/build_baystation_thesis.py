"""Build the canonical Project FALCON Bay Station thesis document.

This generator intentionally describes the approved Phase 1 research baseline:
pressure-derived wave estimation and wind sensing on an ESP32 buoy, LoRa to a
shore Bay Station, and SIM/4G/5G Internet backhaul at the Bay Station. Planned
hardware and AI performance are never presented as field-validated results.
"""

from __future__ import annotations

import math
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
THESIS = ROOT / "THESIS DOCUMENTATION"
VISUALS = THESIS / "visuals" / "baystation-v4"
OUTPUT = THESIS / "BayStation.docx"
SYNC_OUTPUT = THESIS / "PROJECT FALCON-01 - V3 Documentation.docx"
LOGO = ROOT / "data" / "falcon-logo.jpg"
DASHBOARD = THESIS / "visuals" / "dashboard-overview.png"

NAVY = "12304A"
TEAL = "0F766E"
BLUE = "2563EB"
LIGHT = "EAF3F6"
MID = "D6E4EA"
TEXT = "1F2937"
MUTED = "4B5563"
RED = "B91C1C"
AMBER = "B45309"
GREEN = "15803D"


def font(size: int, bold: bool = False):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def rounded_box(draw, xy, fill, outline=NAVY, radius=28, width=4):
    draw.rounded_rectangle(xy, radius=radius, fill="#" + fill, outline="#" + outline, width=width)


def centered(draw, xy, text, fnt, fill=TEXT):
    box = draw.multiline_textbbox((0, 0), text, font=fnt, align="center", spacing=7)
    x = xy[0] - (box[2] - box[0]) / 2
    y = xy[1] - (box[3] - box[1]) / 2
    draw.multiline_text((x, y), text, font=fnt, fill="#" + fill, align="center", spacing=7)


def arrow(draw, start, end, color=TEAL, width=8):
    draw.line([start, end], fill="#" + color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 22
    for offset in (2.55, -2.55):
        p = (end[0] + length * math.cos(angle + offset), end[1] + length * math.sin(angle + offset))
        draw.line([end, p], fill="#" + color, width=width)


def canvas(title: str, subtitle: str):
    image = Image.new("RGB", (1800, 1000), "#F4F7F9")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 130), fill="#" + NAVY)
    draw.text((65, 32), title, font=font(42, True), fill="white")
    draw.text((67, 84), subtitle, font=font(22), fill="#D6E4EA")
    return image, draw


def save_architecture():
    image, draw = canvas("PROJECT FALCON — APPROVED DATA PATH", "No mini PC or cellular Internet modem is installed on the buoy")
    labels = [
        ("MARINE INPUTS", "Pressure\nWind"),
        ("BUOY NODE", "ESP32\nvalidation + buffer"),
        ("PRIMARY LINK", "LoRa\nsite-tested radio"),
        ("BAY STATION", "Mini PC\nSQLite + processing + AI"),
        ("OPERATORS", "Local dashboard\nCloud / remote access"),
    ]
    xs = [190, 535, 880, 1225, 1580]
    colors = ["E0F2FE", "DCFCE7", "FEF3C7", "DBEAFE", "F3E8FF"]
    for index, ((heading, body), x, color) in enumerate(zip(labels, xs, colors)):
        rounded_box(draw, (x - 145, 300, x + 145, 585), color)
        centered(draw, (x, 350), heading, font(23, True), NAVY)
        centered(draw, (x, 460), body, font(28, True), TEXT)
        if index < len(xs) - 1:
            arrow(draw, (x + 150, 440), (xs[index + 1] - 150, 440))
    rounded_box(draw, (1035, 680, 1415, 845), "FFF7ED", AMBER)
    centered(draw, (1225, 735), "SIM / 4G / 5G", font(26, True), AMBER)
    centered(draw, (1225, 790), "Bay Station Internet backhaul", font(22), TEXT)
    arrow(draw, (1225, 585), (1225, 675), AMBER)
    draw.text((60, 900), "Failure behavior: LoRa outage → ESP32 buffers records. Internet outage → Bay Station continues local storage, AI, and dashboard.", font=font(24, True), fill="#" + MUTED)
    path = VISUALS / "01-approved-architecture.png"
    image.save(path, quality=95)
    return path


def save_ipo():
    image, draw = canvas("CONCEPTUAL FRAMEWORK — INPUT, PROCESS, OUTPUT", "The study evaluates an integrated prototype, not a replacement for official coastal instruments")
    blocks = [
        ("INPUT", ["Underwater pressure", "Wind speed and direction", "GPS / power / security", "Requirements and references"]),
        ("PROCESS", ["ESP32 acquisition and validation", "LoRa telemetry and buffering", "Pressure-to-wave processing", "AI prediction and quality checks"]),
        ("OUTPUT", ["Estimated wave height", "10-minute AI prediction", "Readable dashboard and alerts", "Accuracy and reliability evidence"]),
    ]
    xs = [325, 900, 1475]
    fills = ["E0F2FE", "DCFCE7", "FEF3C7"]
    for index, ((heading, items), x, fill) in enumerate(zip(blocks, xs, fills)):
        rounded_box(draw, (x - 235, 250, x + 235, 790), fill)
        centered(draw, (x, 315), heading, font(34, True), NAVY)
        y = 410
        for item in items:
            draw.ellipse((x - 190, y - 5, x - 172, y + 13), fill="#" + TEAL)
            draw.text((x - 155, y - 14), item, font=font(22), fill="#" + TEXT)
            y += 85
        if index < 2:
            arrow(draw, (x + 240, 520), (xs[index + 1] - 240, 520))
    path = VISUALS / "02-conceptual-framework.png"
    image.save(path, quality=95)
    return path


def save_wave_pipeline():
    image, draw = canvas("PRESSURE-DERIVED WAVE ESTIMATION", "Every output retains timestamp, units, source, validity, and calibration state")
    steps = [
        "Raw pressure\n(kPa)", "Range + stale\nvalidation", "Baseline and\nstatic-pressure removal", "Wave-band\nfiltering", "Depth/response\ncorrection", "Wave statistic\n(e.g., Hs)", "Reference\nvalidation"
    ]
    xs = [155, 400, 645, 890, 1135, 1380, 1625]
    for i, (step, x) in enumerate(zip(steps, xs)):
        rounded_box(draw, (x - 105, 330, x + 105, 590), "FFFFFF", TEAL)
        centered(draw, (x, 455), step, font(22, True), TEXT)
        draw.ellipse((x - 25, 650, x + 25, 700), fill="#" + TEAL)
        centered(draw, (x, 675), str(i + 1), font(22, True), "FFFFFF")
        if i < len(xs) - 1:
            arrow(draw, (x + 110, 460), (xs[i + 1] - 110, 460), BLUE, 6)
    rounded_box(draw, (350, 785, 1450, 900), "FEF2F2", RED)
    centered(draw, (900, 842), "Until comparison testing is complete: display ESTIMATED and CALIBRATION REQUIRED — never claim instrument-grade accuracy.", font(24, True), RED)
    path = VISUALS / "03-pressure-wave-pipeline.png"
    image.save(path, quality=95)
    return path


def save_evaluation():
    image, draw = canvas("WHOLE-SYSTEM VALIDATION FRAMEWORK", "Evidence must come from controlled tests; simulator output is presentation data only")
    center = (900, 520)
    rounded_box(draw, (700, 400, 1100, 640), "DBEAFE", BLUE)
    centered(draw, center, "PROJECT FALCON\nVALIDATED PROTOTYPE", font(30, True), NAVY)
    items = [
        ("SENSOR ACCURACY", "reference comparison\nMAE · RMSE · bias", (300, 260)),
        ("COMMUNICATION", "packet delivery · latency\nreconnect · buffering", (900, 240)),
        ("POWER", "measured Wh/day\nautonomy · recovery", (1500, 260)),
        ("MECHANICAL", "stability · ingress\ncorrosion · retrieval", (300, 760)),
        ("USABILITY", "task completion\nreadability · errors", (900, 790)),
        ("AI", "chronological holdout\nbaseline comparison", (1500, 760)),
    ]
    for heading, body, (x, y) in items:
        rounded_box(draw, (x - 190, y - 105, x + 190, y + 105), "FFFFFF", TEAL)
        centered(draw, (x, y - 35), heading, font(23, True), NAVY)
        centered(draw, (x, y + 35), body, font(20), MUTED)
        start_x = x + (190 if x < 900 else -190 if x > 900 else 0)
        start_y = y + (0 if x != 900 else 105 if y < 520 else -105)
        target_x = 700 if x < 900 else 1100 if x > 900 else 900
        target_y = 520 if x != 900 else 400 if y < 520 else 640
        arrow(draw, (start_x, start_y), (target_x, target_y), MID, 5)
    path = VISUALS / "04-validation-framework.png"
    image.save(path, quality=95)
    return path


def save_power_boundary():
    image, draw = canvas("POWER AND DEPLOYMENT BOUNDARIES", "Buoy autonomy and shore Bay Station power are calculated separately")
    rounded_box(draw, (110, 250, 820, 810), "E0F2FE", BLUE)
    centered(draw, (465, 310), "SOLAR-POWERED BUOY", font(32, True), NAVY)
    for y, label in [(410, "Solar panel + charge controller"), (515, "LiFePO₄ battery + protection"), (620, "ESP32 + sensors + LoRa"), (725, "Measured daily energy budget")]:
        rounded_box(draw, (220, y - 42, 710, y + 42), "FFFFFF", TEAL, 18, 3)
        centered(draw, (465, y), label, font(22, True), TEXT)
    rounded_box(draw, (980, 250, 1690, 810), "DCFCE7", GREEN)
    centered(draw, (1335, 310), "SHORE BAY STATION", font(32, True), NAVY)
    for y, label in [(410, "Facility power / separate UPS"), (515, "LoRa gateway + mini PC"), (620, "SIM/4G/5G modem/router"), (725, "Local operation during Internet loss")]:
        rounded_box(draw, (1090, y - 42, 1580, y + 42), "FFFFFF", GREEN, 18, 3)
        centered(draw, (1335, y), label, font(22, True), TEXT)
    arrow(draw, (820, 530), (980, 530), AMBER, 8)
    centered(draw, (900, 475), "LoRa", font(24, True), AMBER)
    path = VISUALS / "05-power-boundaries.png"
    image.save(path, quality=95)
    return path


def generate_visuals():
    VISUALS.mkdir(parents=True, exist_ok=True)
    return [save_architecture(), save_ipo(), save_wave_pipeline(), save_evaluation(), save_power_boundary()]


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, value=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def configure(document):
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(25)
    section.right_margin = Mm(20)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in [("Title", 23, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 13, TEAL), ("Heading 3", 11, BLUE)]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.text = "PROJECT FALCON-01  |  BAY STATION THESIS DOCUMENTATION"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Fullbright College  •  Undergraduate Research Prototype  •  ")
    add_field(footer, "PAGE")
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def p(document, text="", *, bold_lead=None, center=False, italic=False, indent=True):
    para = document.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent and not center:
        para.paragraph_format.first_line_indent = Mm(8)
    if bold_lead and text.startswith(bold_lead):
        lead = para.add_run(bold_lead)
        lead.bold = True
        para.add_run(text[len(bold_lead):])
    else:
        run = para.add_run(text)
        run.italic = italic
    return para


def bullets(document, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        para = document.add_paragraph(item, style=style)
        para.paragraph_format.left_indent = Mm(8)
        para.paragraph_format.first_line_indent = Mm(0)


def table(document, headers, rows, widths=None):
    result = document.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.style = "Table Grid"
    result.autofit = True
    for i, value in enumerate(headers):
        cell = result.rows[0].cells[i]
        cell.text = value
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_index, row in enumerate(rows):
        cells = result.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if row_index % 2:
                shade(cells[i], "F4F7F9")
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs:
                    run.font.size = Pt(8.5)
    return result


def figure(document, path, caption, width=6.6):
    para = document.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(str(path), width=Inches(width))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    run = cap.add_run(caption)
    run.bold = True
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def page_break(document):
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def heading(document, text, level=1):
    return document.add_heading(text, level=level)


def build():
    visuals = generate_visuals()
    document = Document()
    configure(document)

    if LOGO.exists():
        logo_p = document.add_paragraph()
        logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_p.add_run().add_picture(str(LOGO), width=Inches(1.35))
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    run = title.add_run("PROJECT FALCON-01")
    run.bold = True
    run.font.size = Pt(25)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Design and Development of a Solar-Powered Smart Coastal Observation Buoy with a Shore-Based Bay Station for Pressure-Derived Wave Monitoring and AI-Assisted Short-Term Prediction")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    p(document, "Expanded Canonical Thesis Documentation — Version 4.1", center=True, indent=False)
    p(document, "Project FALCON Research Group\nBachelor of Science in Information Technology\nFullbright College", center=True, indent=False)
    p(document, date.today().strftime("%d %B %Y"), center=True, indent=False)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(36)
    r = note.add_run("STATUS: UNDERGRADUATE RESEARCH PROTOTYPE — HARDWARE INTEGRATION AND FIELD VALIDATION PENDING")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(RED)
    page_break(document)

    heading(document, "Document Control")
    table(document, ["Field", "Approved value"], [
        ["Canonical file", "THESIS DOCUMENTATION/BayStation.docx"],
        ["Architecture", "ESP32 buoy → LoRa → shore Bay Station → SIM/4G/5G Internet backhaul"],
        ["Primary measurements", "Pressure-derived estimated wave height; wind speed and direction"],
        ["Pressure candidate", "Holykell HPT604 Type A, provisional 0–2 mH2O vented gauge, 4–20 mA; procurement and validation pending"],
        ["Supporting telemetry", "GPS/time, battery/solar, security, enclosure/system health"],
        ["AI status", "Required shore-based prediction feature; current model is an unvalidated research baseline"],
        ["Physical status", "Prototype geometry under redesign; no fabrication or deployment release"],
        ["Supersedes", "Conflicting onboard Orange Pi, on-buoy cellular, BNO085, load-cell, salinity, and water-temperature descriptions"],
    ])
    heading(document, "Abstract")
    p(document, "Project FALCON-01 proposes an affordable, modular, solar-powered coastal observation buoy for controlled near-shore monitoring. The buoy uses an ESP32 to acquire underwater pressure and wind measurements and to attach supporting GPS, power, security, and health telemetry. Compact packets are transmitted through a site-tested LoRa link to a shore-based Bay Station located at a barangay facility. The Bay Station stores original and derived records in SQLite, estimates wave height from calibrated pressure variation, runs a short-term AI-assisted prediction pipeline, serves a readable four-page dashboard, and uses SIM/4G/5G only as its Internet backhaul. The study evaluates the prototype through sensor-reference comparison, communication reliability, power autonomy, usability, security false-alert behavior, mechanical readiness, and chronological AI validation. The system is an educational research and decision-support prototype; it does not replace PAGASA products, certified oceanographic instruments, navigation systems, or emergency warnings.")
    p(document, "Keywords: coastal observation buoy, ESP32, LoRa, Bay Station, underwater pressure, estimated wave height, wind monitoring, edge computing, short-term prediction", italic=True, indent=False)
    heading(document, "Table of Contents")
    table(document, ["Section", "Page"], [
        ["Document Control and Abstract", "2"],
        ["Chapter 1 — The Problem and Its Background", "4"],
        ["Chapter 2 — Review of Related Literature and System Basis", "8"],
        ["Chapter 3 — Methodology", "10"],
        ["Chapter 4 — Current Implementation Status and Development Plan", "18"],
        ["References", "20"],
        ["Appendices", "21"],
    ])
    heading(document, "List of Figures", 2)
    bullets(document, [
        "Figure 1. Input–Process–Output conceptual framework",
        "Figure 2. Approved Project FALCON data and communication architecture",
        "Figure 3. Pressure-derived wave-estimation pipeline",
        "Figure 4. Existing dashboard software prototype",
        "Figure 5. Buoy and shore Bay Station power boundaries",
        "Figure 6. Whole-system evaluation framework",
    ])
    page_break(document)

    heading(document, "CHAPTER 1 — THE PROBLEM AND ITS BACKGROUND")
    heading(document, "1.1 Background of the Study", 2)
    p(document, "The Philippines depends on coastal waters for fisheries, transportation, tourism, education, environmental research, and community livelihoods. Broad marine forecasts are essential, but conditions at a particular near-shore site can differ from regional descriptions. Smaller schools and local organizations may also lack access to expensive professional instruments, proprietary services, and specialized maintenance.")
    p(document, "Affordable controllers, low-power radios, renewable-energy components, and open-source software allow researchers to investigate a locally maintainable observation platform. Project FALCON combines a solar buoy sensing node with a shore Bay Station. Its purpose is to collect traceable measurements, preserve data during connectivity interruptions, and present current conditions in language that non-technical operators can understand.")
    heading(document, "1.2 Statement of the Problem", 2)
    p(document, "Low-cost monitoring prototypes often demonstrate individual sensors or dashboards without evaluating measurement quality, communication loss, power autonomy, maintenance, and user comprehension as one system. Project FALCON addresses the integration problem while preserving honest distinctions among measured, estimated, predicted, simulated, stale, and unavailable data.")
    bullets(document, [
        "How accurately and repeatably can calibrated underwater-pressure variation estimate the selected wave-height statistic during controlled trials?",
        "How reliably can the ESP32–LoRa–Bay Station path deliver timestamped records and recover from link or Internet outages without silent data loss or duplication?",
        "Can the solar-battery subsystem sustain the required buoy operating schedule under measured load and reduced-sun conditions?",
        "How effectively can persistent geofence and enclosure/tamper rules identify abnormal events without treating ordinary motion or GPS scatter as theft?",
        "Can older or non-technical operators understand the four-page dashboard and complete essential monitoring tasks with minimal assistance?",
        "Does the selected AI method improve a non-AI baseline for one defined short-term wave-height target on chronologically held-out calibrated data?",
    ], numbered=True)
    heading(document, "1.3 General Objective", 2)
    p(document, "To design, develop, and evaluate an affordable solar-powered coastal observation buoy and shore-based Bay Station that provide pressure-derived wave monitoring, wind observations, resilient telemetry, operator-friendly visualization, and AI-assisted short-term wave-height prediction.")
    heading(document, "1.4 Specific Objectives", 2)
    bullets(document, [
        "Integrate pressure and wind sensors with an ESP32 and clearly identified supporting telemetry.",
        "Implement timestamping, validation, calibration metadata, LoRa transmission, and outage buffering.",
        "Develop a documented pressure-processing method that reports an explicitly labeled estimated wave-height statistic.",
        "Implement Bay Station ingestion, SQLite storage, duplicate prevention, alerts, API services, and four-page dashboard hosting.",
        "Develop one short-term AI prediction target and compare it with persistence and moving-average baselines.",
        "Evaluate measurement accuracy, communication, power, usability, security behavior, and controlled deployment readiness.",
    ], numbered=True)
    heading(document, "1.5 Significance of the Study", 2)
    table(document, ["Beneficiary", "Potential value"], [
        ["Coastal communities and local offices", "A locally viewable research tool that may improve awareness of measured site conditions after validation."],
        ["Students and researchers", "A reproducible platform for embedded systems, IoT, databases, visualization, signal processing, and AI evaluation."],
        ["Fullbright College", "An interdisciplinary test platform and documented dataset for future controlled studies."],
        ["Future developers", "Open architecture, interfaces, calibration records, and limitations that can support later improvements."],
    ])
    heading(document, "1.6 Scope and Delimitations", 2)
    p(document, "Phase 1 covers one controlled near-shore prototype, passive single-anchor mooring, an ESP32 buoy node, pressure-derived wave estimation, wind speed/direction, supporting GPS/power/security telemetry, LoRa transport to one shore Bay Station, local storage, one dashboard, alerts, and one short-term prediction target. The recommended pressure candidate is a Holykell HPT604 Type A with provisional 0–2 mH2O vented-gauge range and 4–20 mA output. Its exact order code and continuous-seawater suitability remain approval gates together with the site, LoRa radio/gateway, Bay Station computer, cellular provider, antenna, and cloud endpoint.")
    p(document, "The study excludes official storm, typhoon, tsunami, navigation, or emergency-warning claims; laboratory-grade water quality; salinity; required water temperature; BNO085-based wave measurement; anchor-chain load sensing; autonomous navigation; camera AI; satellite communication; and multi-buoy networking. Simulator and software-test results do not prove physical accuracy or coastal readiness.")
    heading(document, "1.7 Definition of Operational Terms", 2)
    table(document, ["Term", "Operational meaning in this study"], [
        ["Measured", "A direct sensor reading retained with units, source, timestamp, and validity."],
        ["Estimated wave height", "A value derived from underwater pressure using a documented calibration and processing method."],
        ["AI prediction", "A future wave-height estimate produced by a versioned model and never presented as a measured value."],
        ["Near-real-time", "A monitored delivery interval whose numeric target must be frozen and measured during communication testing."],
        ["Bay Station", "The shore computer and gateway that store, process, predict, alert, and serve the dashboard."],
        ["LoRa", "The primary buoy-to-shore radio link; it is not the buoy's direct Internet connection."],
        ["Supporting telemetry", "GPS, time, power, security, and health information required to interpret and operate the system."],
    ])
    figure(document, visuals[1], "Figure 1. Input–Process–Output conceptual framework for the approved Phase 1 study.")
    page_break(document)

    heading(document, "CHAPTER 2 — REVIEW OF RELATED LITERATURE AND SYSTEM BASIS")
    heading(document, "2.1 Low-Cost Coastal Monitoring", 2)
    p(document, "Albaladejo et al. (2012) demonstrated that low-cost shallow-marine monitoring requires the buoy, sensing, wireless communication, and energy supply to be engineered as one platform. Williams et al. (2025) later reinforced the feasibility of compact modular autonomous buoys. These studies support FALCON's integrated approach, but they also show that affordability and sensor assembly alone are not sufficient novelty.")
    heading(document, "2.2 Pressure-Derived Wave Observation", 2)
    p(document, "Bishop and Donelan (1987) established that subsurface pressure can support wave measurement only when attenuation, depth, and frequency response are considered. Bonneton et al. (2018) showed that reconstruction quality also depends on dispersion and nonlinearity. Ardhuin et al. (2019) emphasized that every observation platform and processing method has a response and uncertainty. FALCON must therefore define one wave statistic, document sampling and filtering, retain raw pressure, and compare results against an independent time-aligned reference.")
    heading(document, "2.3 Communication, Power, and Data Quality", 2)
    p(document, "Dreyer et al. (2026) provides emerging evidence that LoRa can support coastal buoy telemetry, but the work is a preprint and cannot substitute for a site survey. Cho et al. (2021) demonstrates why marine sensing, communication, and battery behavior must be evaluated together. Skålvik et al. (2023) identifies calibration drift, corrosion, power limits, intermittent communication, and limited references as recurring threats to ocean-sensor data quality. These findings support FALCON's buffering, quality flags, measured energy budget, and controlled maintenance plan.")
    heading(document, "2.4 Position, Security, and False Alerts", 2)
    p(document, "Knight et al. (2021) demonstrates that low-cost GNSS observations can be useful when accompanied by appropriate processing and reference comparison. For a moored buoy, position changes may result from receiver scatter, anchor swing, currents, tides, maintenance, drag, or theft. FALCON therefore treats GPS as supporting security telemetry and uses persistence, fix quality, and maintenance state rather than declaring theft from a single coordinate.")
    heading(document, "2.5 Short-Term Wave Prediction", 2)
    p(document, "Fan et al. (2020) and Song et al. (2023) demonstrate that sequence models can learn nonlinear wave-height patterns when adequate data are available. Their reported performance cannot be transferred directly to a new low-cost local dataset. FALCON must use chronological partitions, simple baselines, model/version records, input-quality gates, and uncertainty based on observed residuals rather than an arbitrary confidence percentage.")
    heading(document, "2.6 Synthesis and Research Gap", 2)
    p(document, "Prior studies establish the feasibility of low-cost buoys, pressure-based wave reconstruction, renewable power, wireless telemetry, and data-driven prediction separately. The defensible gap is the lack of a reproducible student-scale system evaluated as a whole for a defined local Philippine near-shore use case using pressure-derived wave monitoring, wind observations, a resilient LoRa-to-Bay-Station data path, transparent prediction, and an interface intended for non-technical operators. FALCON claims novelty in locally adapted integration and evidence-based evaluation—not invention of the sensors, LoRa, or AI algorithms.")
    heading(document, "2.7 Conceptual and Technical Basis", 2)
    p(document, "The project follows an input–process–output model. Direct pressure and wind observations enter the ESP32 acquisition layer. Communication and Bay Station services convert validated samples into retained records, derived wave statistics, predictions, and operator messages. Evaluation evidence closes the loop by determining whether each output is accurate, reliable, usable, and safe enough for the stated research scope.")
    page_break(document)

    heading(document, "CHAPTER 3 — METHODOLOGY")
    heading(document, "3.1 Research and Development Design", 2)
    p(document, "The study uses a developmental research approach with iterative requirements analysis, prototype design, implementation, controlled testing, evaluation, and revision. Quantitative measurements are used for sensor accuracy, communication, power, alert performance, usability tasks, and prediction error. Qualitative comments may be collected from authorized evaluators to identify clarity and maintenance problems.")
    heading(document, "3.2 Approved System Architecture", 2)
    figure(document, visuals[0], "Figure 2. Approved Project FALCON data and communication architecture.")
    table(document, ["Layer", "Responsibility", "Failure behavior"], [
        ["Sensors and interfaces", "Produce primary pressure/wind signals and supporting telemetry.", "Report invalid or unavailable states; do not substitute zero."],
        ["ESP32 buoy node", "Acquire, timestamp, validate, frame, secure, transmit, and buffer telemetry.", "Continues sensing and local security during LoRa loss."],
        ["LoRa link", "Carry compact packets to the barangay-hall gateway.", "Triggers buffering and later timestamp-preserving retransmission."],
        ["Bay Station", "Authenticate, deduplicate, store, process, predict, alert, and serve the dashboard.", "Continues local operation during Internet loss."],
        ["Internet backhaul", "Cloud synchronization and authorized remote access only.", "Queues upload while local services remain available."],
    ])
    heading(document, "3.3 Hardware Baseline", 2)
    table(document, ["Function", "Current selection/status", "Required evidence"], [
        ["Pressure", "HPT604 Type A 0–2 mH2O, 4–20 mA deployment candidate; Bar02 bench-only", "Supplier seawater confirmation, exact order code, loop interface, depth comparison, baseline, dynamic response, vent, drift, fouling and endurance"],
        ["Wind", "SparkFun SEN-15901 prototype", "Speed reference, vane alignment, startup threshold, corrosion and cable tests"],
        ["GPS", "Adafruit Ultimate GPS PID 746 prototype", "Stationary scatter, cold/warm start, geofence false-alert tests"],
        ["Power monitoring", "Two INA260 channels", "DMM comparison, polarity, range, connector heating"],
        ["Controller", "ESP32 DevKit", "Pinout, current, reset/recovery, watchdog and buffer behavior"],
        ["LoRa", "Exact radio, gateway, antenna, band, and protocol TBD", "Legal-band review, coverage survey, packet loss and recovery"],
        ["Bay Station", "Development laptop substitute; final mini PC TBD", "OS, startup, storage, cooling, UPS, measured load"],
    ])
    p(document, "Important pressure-sensor gate. The HPT604 candidate uses a protected 12 V 4–20 mA loop, a 150 ohm precision shunt, input protection/filtering, and a 3.3 V ADS1115 receiver. The theoretical loop load is 0.048–0.240 W at 12 V before conversion loss. Its vent tube must terminate in a dry breathable supplier-approved desiccant/breather arrangement. Written confirmation of continuous saltwater compatibility, exact wetted materials, seals, cable and order code is required before purchase. The current Bar02 I2C PCB/connector is incompatible and remains bench-only because the manufacturer requires daily drying and limits continuous immersion.")
    heading(document, "3.4 Data Acquisition and Telemetry", 2)
    bullets(document, [
        "Assign every record a station ID, packet ID/sequence, original sample timestamp, firmware/schema version, source, and quality flags.",
        "Acquire pressure at a sampling rate justified by the target wave band; acquire slower supporting channels at independent documented rates.",
        "Reject or flag out-of-range, stale, disconnected, implausible, and stuck values without silently replacing them.",
        "Buffer a defined duration of records during LoRa outages and preserve original timestamps during retransmission.",
        "Authenticate the buoy identity, prevent duplicate database insertion, and log rejected or malformed packets.",
    ])
    heading(document, "3.5 Pressure-to-Wave Processing", 2)
    figure(document, visuals[2], "Figure 3. Required traceable processing stages for pressure-derived wave estimation.")
    p(document, "Before implementation is frozen, the researchers must define whether the primary output is instantaneous peak-to-trough height, mean height, or significant wave height (Hs). The recommended research output is estimated significant wave height over a stated observation window because the term can be defined, aggregated, compared, and predicted consistently. The team must then freeze sampling frequency, window duration, density assumption, static/tidal removal, filter band, depth-response correction, gap policy, and reference method.")
    p(document, "A simple hydrostatic relationship may support depth conversion, but dynamic surface-wave reconstruction requires documented treatment of sensor depth and frequency-dependent attenuation. All algorithm coefficients and revisions must be versioned. When calibration is incomplete, the dashboard must show ESTIMATED and CALIBRATION REQUIRED.")
    heading(document, "3.6 AI Prediction Method", 2)
    bullets(document, [
        "Freeze one primary target: recommended 10-minute-ahead estimated Hs, subject to adviser and stakeholder approval.",
        "Build persistence and moving-average baselines before testing linear, tree-based, GRU, or LSTM candidates.",
        "Split data chronologically by deployment block or day; never randomly mix adjacent samples across training and test sets.",
        "Report MAE, RMSE, bias, skill relative to persistence, inference time, model size, missing-input behavior, and error by sea-condition range.",
        "Withhold the prediction when pressure input is invalid, stale, or uncalibrated; monitoring continues if AI fails.",
    ])
    heading(document, "3.7 Dashboard and Usability Method", 2)
    p(document, "The approved navigation contains Overview, Buoy Motion, Sensors, and Logs & Alerts. GPS is grouped under Supporting Telemetry rather than becoming a fifth page. Settings remains a compact header control. The interface uses large text, plain labels, clear source/freshness states, and minimal actions for older or non-technical users.")
    if DASHBOARD.exists():
        figure(document, DASHBOARD, "Figure 4. Existing dashboard software prototype; values shown in simulator mode are not field measurements.", width=6.35)
    bullets(document, [
        "Give evaluators realistic tasks: identify current wave estimate, verify GPS/security state, find a stale sensor, acknowledge an alert, and export records.",
        "Record task completion, completion time, navigation errors, assistance required, and label comprehension.",
        "Collect readability and confidence ratings using an adviser-approved usability instrument and sample.",
    ])
    heading(document, "3.8 Security, Privacy, and Data Management", 2)
    p(document, "Physical security uses SECURE, WARNING, ALERT, and DISARMED states. The final implementation must document geofence radius, GPS-quality limits, persistence time, maintenance controls, and enclosure/tamper debounce. Ordinary buoy movement cannot be treated as theft by itself.")
    p(document, "Cybersecurity requirements include per-device identity, secret provisioning outside the repository, packet replay/duplicate protection, least-privilege operator access, TLS for Internet traffic, firewall configuration, backup/export, retention, and audit records. Exact controls remain selection gates. GPS locations must be disclosed only to authorized users when a deployment partner considers them sensitive.")
    heading(document, "3.9 Power Method", 2)
    figure(document, visuals[4], "Figure 5. Separate buoy and shore Bay Station power boundaries.")
    p(document, "The preliminary planning example uses a 12.8 V, 20 Ah battery (256 Wh nominal), 80% usable energy (204.8 Wh), and a 4 W average buoy load (96 Wh/day), yielding approximately 51.2 hours without charging. A 72-hour target at 4 W requires approximately 360 Wh nominal at 80% usable capacity, or 28.1 Ah at 12.8 V; a nominal 30 Ah class battery would provide only a planning margin. A 60 W panel with four peak-sun-hours and 70% net efficiency yields approximately 168 Wh/day. These figures are calculations—not measurements—and must be replaced by logged current and charging data after the exact LoRa and hardware selections are frozen.")
    heading(document, "3.10 Mechanical, Environmental, and Deployment Method", 2)
    p(document, "The replacement physical design must pass buoyancy, freeboard, center-of-gravity, stability/righting, tower load, mooring, corrosion, cable strain, ingress, condensation, sensor exposure, antenna clearance, retrieval, and serviceability review. Existing CAD and blueprint dimensions are references only while the geometry is under redesign. Controlled water tests must begin at an accessible supervised site with permits, deployment limits, retrieval equipment, and a time-aligned reference measurement.")
    heading(document, "3.11 Evaluation Matrix", 2)
    figure(document, visuals[3], "Figure 6. Whole-system evaluation framework.")
    table(document, ["Area", "Primary metrics/evidence", "Release condition"], [
        ["Pressure/wave", "MAE, RMSE, bias, repeatability, gaps, comparison plots", "Defined wave statistic and accepted reference performance"],
        ["Wind", "Reference error, direction code, low-speed startup, repeatability", "Acceptance limits approved and met"],
        ["LoRa/backhaul", "Packet delivery, latency, range, reconnect, duplicates, backlog recovery", "Site path works under expected conditions"],
        ["Power", "Average/peak load, Wh/day, autonomy, reduced-sun recovery", "Measured margin meets deployment target"],
        ["Security", "True/false alerts, detection delay, GPS scatter", "False-alert behavior acceptable"],
        ["Dashboard", "Task completion, errors, time, comprehension and readability", "Adviser-approved usability threshold met"],
        ["AI", "MAE/RMSE/bias, skill vs persistence, latency and availability", "Improvement on untouched chronological test data"],
        ["Mechanical", "Stability, ingress, corrosion control, retrieval inspection", "No unsafe or unresolved critical failure"],
    ])
    heading(document, "3.12 Test Scenarios", 2)
    table(document, ["ID", "Scenario", "Expected observable behavior"], [
        ["P-01", "Pressure disconnected or implausible", "Wave estimate and AI withheld; fault logged; no fabricated zero"],
        ["C-01", "LoRa link interrupted", "ESP32 buffers timestamped packets and retransmits without duplicate insertion"],
        ["I-01", "Bay Station Internet interrupted", "Local ingestion, storage, prediction, dashboard, and alerts continue"],
        ["G-01", "GPS scatter near geofence boundary", "WARNING persistence prevents immediate false ALERT"],
        ["T-01", "Enclosure opened while armed", "Persistent event produces logged ALERT; maintenance DISARMED state remains distinct"],
        ["A-01", "AI model unavailable", "Measured/estimated monitoring remains operational; prediction shows unavailable"],
        ["B-01", "Low battery / reduced sunlight", "Power warning, safe operating policy, and recovery are logged"],
        ["R-01", "Bay Station restart", "Services start automatically and stored records remain intact"],
    ])
    heading(document, "3.13 Ethical and Safety Considerations", 2)
    p(document, "Deployment requires permission from the responsible site authority, navigational and environmental review, safe battery and electrical handling, weather limits, retrieval planning, and clear prototype labeling. Human usability participants require informed consent and only the minimum necessary data. FALCON must never issue or imply official warnings and must direct operators to PAGASA and competent authorities for safety decisions.")
    page_break(document)

    heading(document, "CHAPTER 4 — CURRENT IMPLEMENTATION STATUS AND DEVELOPMENT PLAN")
    heading(document, "4.1 Implemented Software", 2)
    bullets(document, [
        "ESP32 PlatformIO firmware shell, diagnostics, and versioned telemetry framing prototype.",
        "Python Bay Station service with simulator/bench ingestion, SQLite, alerts, grouped API, static dashboard hosting, and prediction baseline.",
        "Responsive dashboard and explicit LIVE, SIMULATED, ESTIMATED, CALIBRATION REQUIRED, STALE, and OFFLINE states.",
        "Automated backend tests and dashboard production build.",
    ])
    heading(document, "4.2 Not Yet Validated", 2)
    bullets(document, [
        "Procurement, revised 4–20 mA interface, continuous-saltwater confirmation, and physical calibration coefficients for the HPT604 candidate.",
        "Exact LoRa module, legal band, gateway, antennas, packet protocol, and site coverage.",
        "Final Bay Station mini PC, OS image, modem/provider, UPS, cloud endpoint, and cybersecurity controls.",
        "Final mechanical geometry, waterproofing, corrosion protection, mooring, autonomy, and coastal endurance.",
        "Field-trained AI accuracy and stakeholder usability results.",
    ])
    heading(document, "4.3 Risk Register", 2)
    table(document, ["Risk", "Impact", "Required mitigation"], [
        ["HPT604 exact seawater configuration unconfirmed", "High", "Obtain written supplier confirmation; validate vent, materials, cable, drift, fouling and staged endurance before unattended use."],
        ["Existing PCB targets Bar02 I2C", "High", "Do not fabricate; redesign and bench-test the protected 4–20 mA receiver and keyed connector."],
        ["Undefined LoRa path", "High", "Freeze legal band/hardware and complete a site coverage survey before PCB/deployment release."],
        ["Insufficient calibrated AI data", "High", "Use transparent baseline; collect data; avoid accuracy claims until held-out evaluation."],
        ["Power assumptions differ from measured load", "High", "Log all operating modes and redesign capacity from measured Wh/day."],
        ["Water ingress/corrosion", "High", "Ingress and salt-exposure tests, inspection schedule, strain relief, and serviceable seals."],
        ["GPS false security alerts", "Medium", "Measure scatter and use fix quality, persistence, and maintenance states."],
        ["Operator confusion", "Medium", "Four-page interface, plain labels, task-based testing, and iterative revision."],
    ])
    heading(document, "4.4 Milestone Plan", 2)
    table(document, ["Milestone", "Deliverable", "Exit evidence"], [
        ["M1 Scope freeze", "Signed architecture, site/use case, wave statistic, and AI target", "Adviser approval record"],
        ["M2 Component freeze", "Exact BOM, datasheets, pinout, LoRa/Bay Station selection", "Selection register complete"],
        ["M3 Bench integration", "Calibrated sensors and packet path", "Bench logs and acceptance sheets"],
        ["M4 Controlled water test", "Wave-reference, wind, security, and power dataset", "Traceable synchronized records"],
        ["M5 Software evaluation", "Dashboard usability and AI baseline/model results", "Metrics and held-out evaluation"],
        ["M6 Final revision", "Thesis, presentation, BOM, diagrams, dataset, and limitations", "Adviser-approved release"],
    ])
    heading(document, "4.5 Expected Outputs", 2)
    bullets(document, [
        "One serviceable controlled-trial coastal buoy prototype and one shore Bay Station prototype.",
        "A traceable dataset containing original pressure, derived wave statistics, wind, supporting telemetry, and quality metadata.",
        "A documented non-AI baseline and one evaluated short-term AI prediction model.",
        "A readable dashboard, alerts, logs, and export workflow.",
        "Calibration, communication, energy, usability, security, mechanical, and limitations records.",
    ])
    page_break(document)

    heading(document, "REFERENCES")
    references = [
        "Albaladejo, C., Soto, F., Torres, R., Sánchez, P., & López, J. A. (2012). A low-cost sensor buoy system for monitoring shallow marine environments. Sensors, 12(7), 9613–9634. https://doi.org/10.3390/s120709613",
        "Ardhuin, F., et al. (2019). Observing sea states. Frontiers in Marine Science, 6, Article 124. https://doi.org/10.3389/fmars.2019.00124",
        "Bishop, C. T., & Donelan, M. A. (1987). Measuring waves with pressure transducers. Coastal Engineering, 11(4), 309–328. https://doi.org/10.1016/0378-3839(87)90031-7",
        "Holykell. (n.d.). HPT604 Type A level sensor datasheet. https://www.holykell.com/wp-content/uploads/2023/08/HPT604A-Level-sensor-Datasheet-Holykell-V26-CS-1.pdf",
        "Blue Robotics. (n.d.). Bar high-resolution depth/pressure sensors guide. https://bluerobotics.com/learn/bar-sensors-guide/",
        "Bonneton, P., Lannes, D., Martins, K., & Michallet, H. (2018). A nonlinear weakly dispersive method for recovering the elevation of irrotational surface waves from pressure measurements. Coastal Engineering, 138, 1–8. https://doi.org/10.1016/j.coastaleng.2018.04.005",
        "Cho, J., et al. (2021). Seawater battery-based wireless marine buoy system with battery degradation prediction and multiple power optimization capabilities. IEEE Access, 9, 104104–104114. https://doi.org/10.1109/ACCESS.2021.3098846",
        "Dreyer, L. W., et al. (2026). OLB: An open LoRa buoy for coastal water measurements [Preprint]. arXiv. https://doi.org/10.48550/arXiv.2601.05615",
        "Fan, S., Xiao, N., & Dong, S. (2020). A novel model to predict significant wave height based on long short-term memory network. Ocean Engineering, 205, Article 107298. https://doi.org/10.1016/j.oceaneng.2020.107298",
        "Knight, P. J., Bird, C. O., Sinclair, A., Higham, J., & Plater, A. J. (2021). Beach deployment of a low-cost GNSS buoy for determining sea-level and wave characteristics. Geosciences, 11(12), Article 494. https://doi.org/10.3390/geosciences11120494",
        "Philippine Atmospheric, Geophysical and Astronomical Services Administration. (n.d.). Weather terminologies. https://www.pagasa.dost.gov.ph/information/weather-terminologies",
        "Skålvik, A. M., Saetre, C., Frøysa, K.-E., Bjørk, R. N., & Tengberg, A. (2023). Challenges, limitations, and measurement strategies to ensure data quality in deep-sea sensors. Frontiers in Marine Science, 10, Article 1152236. https://doi.org/10.3389/fmars.2023.1152236",
        "Song, T., et al. (2023). Prediction of significant wave height based on EEMD and deep learning. Frontiers in Marine Science, 10, Article 1089357. https://doi.org/10.3389/fmars.2023.1089357",
        "Williams, Z., Soto Calvo, M. A., Lee, H. S., Aljber, M., & Jeong, J.-S. (2025). A low-cost autonomous multi-functional buoy for ocean currents and seawater parameter monitoring, and particle tracking. Journal of Marine Science and Engineering, 13(9), Article 1629. https://doi.org/10.3390/jmse13091629",
    ]
    for ref in references:
        para = document.add_paragraph(ref)
        para.paragraph_format.left_indent = Mm(8)
        para.paragraph_format.first_line_indent = Mm(-8)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading(document, "APPENDIX A — REQUIRED RECORDS BEFORE FIELD CLAIMS")
    bullets(document, [
        "Signed scope and architecture freeze", "Exact BOM and archived datasheets", "Pinout, schematic, PCB, wiring, and connector release",
        "Pressure and wind calibration sheets", "LoRa and Bay Station backhaul survey", "Power-load and autonomy worksheet",
        "Ingress, corrosion, stability, and retrieval checklists", "Raw and processed dataset with algorithm/model versions",
        "Dashboard usability forms and summary", "Risk, incident, maintenance, and change-control records",
    ])
    heading(document, "APPENDIX B — THESIS-SAFE CLAIMS")
    table(document, ["Use", "Avoid until validated"], [
        ["Pressure-derived estimated wave height", "Accurate or instrument-grade wave height"],
        ["AI-assisted short-term prediction research feature", "Reliable coastal forecast"],
        ["LoRa is the approved primary path; exact range pending survey", "Guaranteed long-range communication"],
        ["Software prototype passed repository tests", "Complete deployed system is validated"],
        ["Potential local monitoring and research value", "Replacement for PAGASA or official warnings"],
    ])
    heading(document, "APPENDIX C — APPROVAL GATES")
    bullets(document, [
        "Named deployment partner, site, user need, permits, and reference access",
        "Defined significant-wave-height method, sampling rate, window, filters, and acceptance limits",
        "Exact HPT604 order code, continuous-seawater confirmation, revised loop interface, and pressure-sensor validation",
        "Exact LoRa hardware, legal band, coverage, security, buffering, and antenna plan",
        "Final mechanical design and engineering review",
        "Measured energy budget and autonomous operating target",
        "AI dataset sufficiency, chronological evaluation, and model release rule",
    ], numbered=True)

    document.core_properties.title = "Project FALCON-01 — Bay Station Thesis Documentation"
    document.core_properties.subject = "Canonical shore Bay Station, LoRa, pressure-wave, and AI-assisted architecture"
    document.core_properties.author = "Project FALCON Research Group"
    document.core_properties.keywords = "FALCON, coastal buoy, Bay Station, LoRa, pressure-derived wave height, AI"
    document.core_properties.comments = "Generated from the adviser-aligned Project FALCON v8.2 repository baseline. Planned hardware and AI claims require validation."
    document.save(OUTPUT)
    SYNC_OUTPUT.write_bytes(OUTPUT.read_bytes())
    print(OUTPUT)
    print(SYNC_OUTPUT)


if __name__ == "__main__":
    build()
