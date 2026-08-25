"""Build the adviser-revised FALCON thesis DOCX from the latest V2 template."""

from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Mm, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "THESIS DOCUMENTATION" / "PROJECT FALCON-01 -  V2 REVISED.docx"
OUTPUT = ROOT / "THESIS DOCUMENTATION" / "PROJECT FALCON-01 - V3 ADVISER REVISED.docx"


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if not child.tag.endswith("sectPr"):
            body.remove(child)


def heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)


def paragraph(document: Document, text: str, bold_lead: str | None = None) -> None:
    item = document.add_paragraph()
    item.paragraph_format.first_line_indent = Inches(0.5)
    item.paragraph_format.line_spacing = 1.5
    item.paragraph_format.space_after = Pt(6)
    item.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        item.add_run(bold_lead).bold = True
        item.add_run(text[len(bold_lead):])
    else:
        item.add_run(text)


def bullets(document: Document, items: list[str]) -> None:
    for text in items:
        item = document.add_paragraph(f"• {text}")
        item.paragraph_format.left_indent = Inches(0.25)
        item.paragraph_format.line_spacing = 1.5


def numbered(document: Document, items: list[str]) -> None:
    for index, text in enumerate(items, 1):
        item = document.add_paragraph(f"{index}. {text}")
        item.paragraph_format.left_indent = Inches(0.25)
        item.paragraph_format.line_spacing = 1.5


def table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    result = document.add_table(rows=1, cols=len(headers))
    try:
        result.style = "Table Grid"
    except KeyError:
        pass
    for index, value in enumerate(headers):
        result.rows[0].cells[index].text = value
    for row in rows:
        cells = result.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    document = Document(SOURCE)
    clear_body(document)
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("PROJECT FALCON-01").bold = True
    title.runs[0].font.size = Pt(20)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Design and Development of a Solar-Powered Smart Coastal Observation Buoy for Real-Time Coastal Monitoring and Pressure-Based Wave-Height Estimation").bold = True
    subtitle.runs[0].font.size = Pt(15)
    version = document.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version.add_run("V3 — ADVISER REVISED\nUndergraduate Thesis Documentation\n25 August 2026")
    document.add_page_break()

    heading(document, "Executive Summary")
    paragraph(document, "Project FALCON-01 is a low-cost, modular, solar-powered smart coastal observation buoy designed for near-real-time local monitoring. An ESP32 acquires pressure, position, wind, environmental, power, system-health, and security data. The ESP32 sends validated telemetry through USB serial/UART to an Orange Pi Zero 3, which stores records locally, exposes a REST API, and serves a responsive browser dashboard without requiring Internet connectivity.")
    paragraph(document, "The adviser-revised Phase 1 study uses a Blue Robotics Bar02 or compatible waterproof pressure sensor as the primary wave input. The device directly measures underwater pressure variation. Software preserves raw pressure, filters noise, establishes a documented baseline, and converts the dynamic pressure component into an estimated wave-height value. The project therefore uses the wording pressure-based estimated wave height and does not claim that the pressure sensor directly measures laboratory-grade wave height.")
    paragraph(document, "The required BNO085 IMU and anchor-chain load-cell/HX711 concepts have been removed from the primary design to keep the undergraduate scope achievable. The buoy uses passive mooring with adequate line scope for tides and normal wave motion. Security is provided by persistent GPS geofence monitoring, a generic vibration/tamper input, an enclosure reed or limit switch, and a buzzer. Artificial intelligence is optional and supporting; all core monitoring, logging, security, and visualization functions operate without it.")

    heading(document, "1. Project Background")
    paragraph(document, "The Philippines has extensive coastlines that support fisheries, transport, tourism, education, environmental research, and community livelihoods. Localized coastal conditions can differ from broader forecasts, yet continuous observation equipment may be difficult for small schools and communities to acquire and maintain. Commercial oceanographic buoys provide professional measurements but can involve high cost, proprietary systems, specialized servicing, and infrastructure beyond the resources of an undergraduate project.")
    paragraph(document, "Affordable microcontrollers, single-board computers, solar components, and open-source software make it possible to investigate a smaller local-first platform. FALCON combines deterministic sensor acquisition on an ESP32 with local storage and visualization on an Orange Pi. Its purpose is educational and research-oriented monitoring, not replacement of professional oceanographic instruments or official agencies.")

    heading(document, "2. Statement of the Problem")
    paragraph(document, "Many low-cost monitoring prototypes demonstrate sensors and dashboards but do not provide a complete serviceable platform that combines traceable pressure-based wave estimation, local environmental data, power autonomy, local data retention, and basic anti-theft/tamper awareness. The study asks whether these functions can be integrated and evaluated within an affordable undergraduate prototype while clearly distinguishing measured, estimated, simulated, stale, and uncalibrated data.")
    numbered(document, [
        "How can a low-cost solar-powered buoy acquire and retain near-real-time coastal measurements using an ESP32 and local Orange Pi edge computer?",
        "How accurately and repeatably can calibrated underwater-pressure variation be processed into an estimated wave-height signal under controlled conditions?",
        "How reliably can GPS geofence, vibration/tamper, and enclosure-access rules detect persistent security events without being triggered by normal wave movement?",
        "How clearly can a four-page local dashboard communicate wave, environment, GPS, power, security, health, and alert information?",
        "How does the prototype perform in sensor accuracy, communication reliability, dashboard usability, energy use, data retention, and system recovery?",
        "If optional AI is evaluated, does it improve a documented baseline without interrupting the core monitoring system?",
    ])

    heading(document, "3. Research Gap and Proposed Novelty")
    paragraph(document, "Existing low-cost IoT studies frequently emphasize acquisition and display, while professional platforms emphasize calibrated accuracy at substantially greater cost. A practical gap remains for an open, modular, locally operated educational prototype that integrates pressure-based wave estimation, environmental context, solar power, local logging, serviceability, and simple security monitoring for controlled Philippine coastal trials.")
    paragraph(document, "FALCON does not claim to invent pressure sensing, GPS geofencing, or wave science. Its proposed novelty is the transparent integration and evaluation of these established technologies in an affordable local-first platform. The study will report limitations, calibration state, and actual test evidence instead of presenting planned hardware or simulator output as completed field performance.")

    heading(document, "4. Objectives")
    heading(document, "4.1 General Objective", 2)
    paragraph(document, "To design, develop, and evaluate a low-cost solar-powered smart coastal observation buoy that provides local real-time coastal monitoring and pressure-based estimated wave height through an ESP32–Orange Pi architecture and responsive web dashboard.")
    heading(document, "4.2 Specific Objectives", 2)
    numbered(document, [
        "Integrate pressure, GPS, wind, environmental, power, system-health, and security channels using documented interfaces and calibration states.",
        "Develop a traceable pipeline that retains raw and filtered pressure, baseline, optional depth, estimated wave height, validity, timestamp, and calibration metadata.",
        "Implement persistent/debounced geofence, tamper, and enclosure security rules with SECURE, WARNING, ALERT, and DISARMED states.",
        "Implement local serial ingestion, SQLite storage, REST API, logs, alerts, and automatic service recovery on the Orange Pi.",
        "Develop a responsive dashboard with Overview, Buoy Motion, Sensors, and Logs & Alerts as primary pages while keeping motion visualization optional.",
        "Evaluate subsystem accuracy, reliability, latency, false alerts, usability, power consumption, and controlled deployment readiness.",
    ])

    heading(document, "5. Scope and Delimitations")
    paragraph(document, "Phase 1 covers a single near-shore prototype, passive single-anchor mooring, local ESP32 acquisition, USB/UART transfer, Orange Pi local processing, pressure-based estimated wave height, core/supporting environmental readings, power monitoring, basic security, local logging, and a browser dashboard. Physical models marked TBD require selection before final wiring or procurement.")
    paragraph(document, "The study does not provide official weather, storm, typhoon, tsunami, navigation, or emergency warnings. It does not claim laboratory-grade salinity, professional oceanographic accuracy, autonomous navigation, satellite communication, camera AI, or multi-buoy operation. Internet connectivity is optional. AI wave prediction is an optional extension and not a required study outcome.")

    heading(document, "6. System Architecture")
    table(document, ["Layer", "Primary responsibility", "Failure behavior"], [
        ["Sensors and interfaces", "Produce raw pressure, GPS, wind, environment, power, and security signals", "Invalid/unavailable values are reported, not replaced by zero"],
        ["ESP32", "Acquire, timestamp, validate, debounce, apply calibration, frame serial telemetry", "Continues acquisition if Orange Pi is unavailable"],
        ["Orange Pi Zero 3", "Ingest, filter, estimate waves, apply geofence rules, store, serve API/dashboard", "Restarts services and preserves local records where possible"],
        ["Dashboard", "Display current state, sensors, logs, alerts, and optional assistant", "Shows OFFLINE/STALE instead of fabricated data"],
        ["Optional AI", "Experimental short-term prediction after baseline validation", "May fail or be disabled without affecting monitoring"],
    ])

    heading(document, "7. Hardware Components")
    table(document, ["Group", "Component", "Purpose and status"], [
        ["Core", "Bar02 or compatible pressure sensor", "Raw pressure and calibrated estimated wave height; selected family"],
        ["Core", "GPS receiver", "Position, time, fix quality, geofence; exact model TBD"],
        ["Core", "Wind speed/direction", "Local wind context; exact models TBD"],
        ["Supporting", "Sealed DS18B20", "Water temperature; reference comparison required"],
        ["Supporting", "Conductivity/salinity interface", "Estimated indicator only; exact model and calibration TBD"],
        ["Health", "Battery, solar, enclosure temperature", "Energy and electronics health; exact interfaces verified before fabrication"],
        ["Security", "GPS geofence, tamper input, enclosure switch, buzzer", "Debounced/persistent anti-theft awareness; exact hardware TBD where stated"],
        ["Controller", "ESP32 DevKit", "Deterministic sensor acquisition and serial telemetry"],
        ["Edge", "Orange Pi Zero 3 (4 GB)", "Local database, API, dashboard, optional AI"],
    ])
    paragraph(document, "The BNO085 remains only as an optional deprecated historical prototype and is not required. Load cell/HX711 mooring-tension sensing is removed. Exact datasheets, logic voltages, connector pinouts, footprints, current requirements, and environmental ratings must be verified before final PCB release.")

    heading(document, "8. Pressure-Based Wave Estimation Method")
    numbered(document, [
        "Read timestamped raw underwater pressure in documented engineering units.",
        "Reject missing, out-of-range, disconnected, or stale samples while retaining diagnostic reasons.",
        "Apply a documented low-pass/band-pass or equivalent filter selected through controlled data analysis.",
        "Establish a pressure baseline at a measured installation depth and record temperature and calibration conditions.",
        "Isolate the dynamic pressure variation associated with water-surface motion.",
        "Apply the documented hydrostatic/dynamic conversion and calibration coefficient to estimate wave height.",
        "Report estimated wave height, quality/validity, timestamp, source, baseline, and calibration state.",
        "Compare results with a documented reference method and report MAE, RMSE, bias, repeatability, and limitations.",
    ])
    paragraph(document, "Until reference testing is complete, displayed values must say ESTIMATED and CALIBRATION REQUIRED. Simulator values must additionally say SIMULATED.")

    heading(document, "9. Security and Tamper Method")
    paragraph(document, "The deployment reference coordinate and geofence radius are configuration values. A position outside the radius enters WARNING while the system evaluates fix quality and persistence. A sustained violation becomes ALERT. Vibration/tamper and enclosure-switch inputs use electrical filtering where required plus firmware debounce/persistence. Ordinary wave movement does not constitute tampering. Authorized maintenance uses DISARMED and creates an event record.")
    table(document, ["State", "Meaning", "Dashboard response"], [
        ["SECURE", "Inside geofence; enclosure closed; no persistent tamper", "Normal status"],
        ["WARNING", "Transient/uncertain event under persistence evaluation", "Review message"],
        ["ALERT", "Persistent geofence/tamper/enclosure event", "Prominent alert and buzzer rule"],
        ["DISARMED", "Authorized maintenance mode", "Logged maintenance indicator"],
    ])

    heading(document, "10. Software, API, and Data")
    paragraph(document, "The ESP32 publishes versioned newline-delimited JSON through USB serial/UART. The Orange Pi validates frames, stores the original payload in SQLite, calculates current derived state, persists alerts/events, and serves the dashboard. The grouped API contains system, wave, environment, GPS, power, security, health, assistant, and alerts sections. Legacy routes remain temporarily available for backward compatibility.")
    bullets(document, [
        "LIVE — connected physical source.", "SIMULATED — presentation simulator.", "ESTIMATED — derived value rather than direct measurement.",
        "CALIBRATION REQUIRED — no completed reference calibration.", "STALE — update age exceeds the configured limit.",
        "OFFLINE — source unavailable.", "OPTIONAL — nonessential feature whose failure cannot stop monitoring.",
    ])

    heading(document, "11. Dashboard Design")
    numbered(document, [
        "Overview: estimated wave and pressure/calibration state, environment, GPS, power, security, health, alerts, and optional FALCON Assistant.",
        "Buoy Motion: optional interactive 3D response model driven by estimated sea context, not a required IMU measurement.",
        "Sensors: grouped core, supporting, health, and security channels with units, quality, source, update age, and calibration status.",
        "Logs & Alerts: current warnings, security/calibration/operator events, persisted telemetry, acknowledgement, search, and export.",
    ])
    paragraph(document, "Settings is a compact header icon. Optional AI is hidden by default. The FALCON Assistant is a rule-based animation with NORMAL, WARNING, ALERT, and OFFLINE states; it is not a chatbot or autonomous decision-maker.")

    heading(document, "12. Testing and Evaluation Plan")
    table(document, ["Test area", "Evidence and metrics"], [
        ["Sensors", "Reference comparison, range, repeatability, invalid/disconnect behavior"],
        ["Wave estimate", "Known reference, MAE, RMSE, bias, repeatability, calibration record"],
        ["Security", "Geofence accuracy, persistence, false positives/negatives, tamper/enclosure tests"],
        ["Communication", "Packet loss, latency, stale detection, reconnect and restart recovery"],
        ["Storage/API", "Retention, timestamp integrity, schema tests, export verification"],
        ["Power", "Normal/peak current, conversion efficiency, autonomy, charge recovery, brownout"],
        ["Dashboard", "Desktop/mobile responsiveness, readability, task completion, label comprehension"],
        ["Mechanical", "Buoyancy, stability, water ingress, corrosion controls, mooring scope and retrieval"],
        ["Optional AI", "Separate baseline comparison and held-out data; no core-system dependency"],
    ])

    heading(document, "13. Risk, Ethics, and Safety")
    paragraph(document, "Marine deployment requires permission, site-risk review, electrical protection, waterproofing, safe battery handling, retrieval planning, and weather limits. GPS data and future camera features require privacy controls. The dashboard must state that FALCON is a research prototype and does not replace PAGASA, coast guard instructions, navigation equipment, or emergency-warning systems.")

    heading(document, "14. Current Implementation and Limitations")
    paragraph(document, "The repository currently includes the ESP32 firmware shell, diagnostic portal, versioned serial frame, Python simulator/edge service, SQLite storage, deterministic alerts, grouped API, four-page dashboard, optional motion visualization, optional assistant, and optional presentation prediction. The software build and automated tests demonstrate implementation behavior only. They do not prove physical sensor accuracy or coastal readiness.")
    paragraph(document, "Physical pressure calibration, exact supporting/security part selection, final PCB/wiring release, Orange Pi installation, waterproofing, power autonomy, and controlled field trials remain pending. These limitations must remain visible in presentations, results, and conclusions.")

    heading(document, "15. Expected Output and Beneficiaries")
    paragraph(document, "The expected output is a documented, serviceable research prototype with traceable coastal readings, pressure-based estimated wave height, local records, security awareness, and a usable dashboard. Potential beneficiaries include students, researchers, educational institutions, coastal communities, and local organizations seeking an accessible platform for controlled observation and future study. Operational adoption requires further validation and coordination with competent authorities.")

    heading(document, "16. Development Roadmap")
    numbered(document, [
        "Approve exact component models and datasheets.", "Freeze the adviser-approved electrical interfaces and revised PCB.",
        "Bench-integrate Bar02, GPS, wind, DS18B20, conductivity, health, and security channels.",
        "Complete pressure/environment/security calibration and controlled reference tests.",
        "Install and harden Orange Pi automatic services.", "Complete enclosure, solar, mooring, and safe controlled water trials.",
        "Analyze results and revise claims based on evidence.", "Evaluate optional AI only if sufficient calibrated data and time remain.",
    ])

    heading(document, "17. Documentation Status")
    paragraph(document, "This V3 adviser-revised document supersedes conflicting V2 descriptions. The repository master context is docs/PROJECT_CONTEXT.md v6.0. Older CAD, Wokwi, motion, IMU, forecast, and PCB records may remain for historical traceability but are not the current required Phase 1 baseline unless revised and explicitly approved.")

    document.add_page_break()
    heading(document, "Appendix A — Approved Telemetry Sections")
    bullets(document, ["system", "wave", "environment", "gps", "power", "security", "health", "assistant", "alerts"])
    heading(document, "Appendix B — Immediate Documentation Records Required")
    bullets(document, [
        "Exact component and supplier register", "Sensor calibration sheets", "Pressure installation-depth and baseline record",
        "GPS geofence configuration and test record", "Tamper/enclosure debounce and false-alert record",
        "Power-load and autonomy worksheet", "Waterproofing and pre-deployment checklist", "Controlled-test dataset and analysis",
    ])

    document.core_properties.title = "Project FALCON-01 - V3 Adviser Revised"
    document.core_properties.subject = "Pressure-based smart coastal observation buoy"
    document.core_properties.comments = "Generated from the latest V2 structural baseline and revised to the 2026-08-25 adviser direction."
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
