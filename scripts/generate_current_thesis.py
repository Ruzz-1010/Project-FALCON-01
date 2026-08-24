"""Revise the original FALCON thesis while preserving its Word formatting."""

from pathlib import Path
from shutil import copy2
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "THESIS DOCUMENTATION" / "PROJECT FALCON-01.docx"
OUTPUT = ROOT / "THESIS DOCUMENTATION" / "PROJECT FALCON-01 - CURRENT V2 REVISED.docx"
VISUALS = ROOT / "THESIS DOCUMENTATION" / "visuals"
SCREENSHOTS = Path(r"C:\Users\Admin\Pictures\Screenshots")

# Paragraph indices intentionally match the original thesis. Replacing text in
# place retains its sections, page setup, styles, headers, footers, and table.
REPLACEMENTS = {
    10: "The current Revision 5 prototype uses a compact single-body marine-grade HDPE main float with a rounded tapered underwater keel, low central ballast, and single-point mooring. A corrosion-resistant frame transfers tower loads to the lower buoy support. Two opposed 30 W solar panels, an MPPT charge controller, and a 12 V LiFePO₄ battery support autonomous operation with practical maintenance access.",
    19: "The current prototype is built around a compact single-body marine-grade HDPE float with a rounded tapered underwater keel, low central ballast, and single-point mooring. Sensor data are acquired by an ESP32, while an Orange Pi Zero 3 (4 GB) is the selected local edge computer. The AI workflow uses validated real-time and historical measurements to estimate wave height over a short 5–15 minute horizon and classify sea conditions as Calm, Moderate, or Rough.",
    83: "Project FALCON combines affordable embedded sensing, two-panel solar power, local edge computing, a maintainable sealed electronics pod, and focused AI-assisted wave analysis in one near-shore research platform.",
    210: "Orange Pi Zero 3 (4 GB) for local edge processing",
    211: "Revision 5 single-body marine-grade HDPE floating platform",
    212: "Rounded tapered underwater keel",
    213: "Corrosion-resistant lower-to-main structural support frame",
    214: "Adjustable low central ballast system",
    215: "Single-point mooring and anchor system",
    219: "BNO085 IMU (SPI)",
    220: "Blue Robotics Bar02 water-pressure sensor (I²C)",
    221: "Wind-speed sensor",
    222: "Wind-direction sensor through ADS1115 ADC",
    223: "GPS module (UART)",
    224: "INA260 battery monitor",
    225: "Second INA260 solar monitor",
    226: "MCP9808 electronics-enclosure temperature sensor",
    227: "Optional sensor",
    228: "Sealed DS18B20 water-temperature sensor",
    255: "Two opposed 30 W solar panels",
    297: "The system supports a focused onboard AI workflow for 5–15 minute wave-height prediction and Calm, Moderate, or Rough sea-condition classification.",
    301: "The current Revision 5 buoy uses a compact single-body HDPE float, rounded tapered underwater keel, low central ballast, single-point mooring, and a corrosion-resistant load path from the lower buoy frame to the tapered equipment tower.",
    322: "Two Opposed 30 W Solar Panels",
    332: "ESP32 Controller              Orange Pi Zero 3 (Edge AI)",
    353: "The Revision 5 mechanical subsystem provides buoyancy, self-righting support, waterproof protection, and a continuous structural load path for the elevated electronics, solar, and sensor assemblies.",
    354: "The current floating platform consists of:", 355: "Revision 5 Main Float",
    356: "Marine-grade HDPE single-body floating structure", 357: "Nominal outside diameter: 650 mm",
    358: "Cylindrical body height: 380 mm", 359: "Rounded tapered underwater keel depth: 240 mm",
    360: "Stability and Structural System",
    361: "Stability is provided by the single-body float geometry, rounded tapered keel, adjustable low central ballast, and controlled single-point mooring configuration.",
    362: "The current arrangement is intended to:", 363: "Lower the center of gravity",
    364: "Support self-righting behavior", 365: "Reduce the deployed footprint",
    366: "Transfer tower loads into the lower buoy frame", 367: "Support the elevated electronics and sensor tower",
    368: "Structural Support Frame", 369: "Material:",
    370: "Marine-grade corrosion-resistant metal, subject to final fabrication selection", 371: "Function:",
    372: "Connect the lower buoy support to the main frame", 373: "Carry tower, pod, solar-panel, and sensor loads",
    374: "Provide corrosion-resistant reinforcement for the marine environment",
    376: "An adjustable ballast positioned below the rounded keel lowers the center of gravity. Its final mass and position require controlled flotation, heel, and roll-recovery testing.",
    378: "A single-point mooring line connects the buoy to its anchor while permitting controlled response to tides and waves. Final deployment requires verified mooring-load and recovery procedures.",
    392: "Edge Computer", 393: "Orange Pi Zero 3 (4 GB), selected; physical integration pending",
    395: "Local AI-assisted wave analysis", 398: "Local dashboard and REST API hosting",
    399: "5–15 minute wave-height prediction and sea-state classification",
    403: "Two Opposed 30 W Solar Panels",
    413: "Sealed Electronics-Pod Thermal System",
    414: "The current design keeps the electronics pod sealed from salt air and spray while circulating internal air across component heat sinks and a rear thermal-transfer path.",
    415: "Internal Recirculation", 416: "Two internal fans move enclosure air across the electronics and heat sinks; they do not exchange outside air directly.",
    417: "Thermal components:", 418: "Horizontal component heat sinks", 419: "Vertical thermal bridge",
    420: "Internal airflow guides and recirculation fans", 421: "Purpose:",
    422: "Move heat from electronic components toward the sealed rear heat-transfer assembly.",
    424: "External Heat Rejection", 425: "The rear external finned heat sink releases conducted heat to ambient air without opening the enclosure.",
    426: "Protection:", 427: "Front maintenance door with raised sealing lip",
    428: "Dual EPDM gasket barriers and compression latches", 429: "Marine-protected cable entries and corrosion-resistant hardware",
    430: "Purpose:", 431: "Maintain serviceability while limiting saltwater ingress; thermal and ingress performance must still be validated by test.",
    434: "Project FALCON uses a focused Phase 1 sensor set directly related to wave measurement, position, wind, power, and enclosure temperature.",
    435: "Wave and Motion Sensors", 436: "BNO085 IMU (SPI)", 437: "Blue Robotics Bar02 water-pressure sensor (I²C)",
    438: "GPS module (UART)", 439: "Optional sealed DS18B20 water-temperature sensor",
    440: "Pressure baseline and motion calibration are required before reporting wave height.",
    442: "Wind Sensors", 443: "Wind-speed sensor", 444: "Wind-direction sensor", 445: "ADS1115 ADC for wind-direction input",
    446: "Wind hardware selection and calibration remain subject to physical validation.",
    447: "The Phase 1 prototype does not include a general humidity or weather-station sensor suite.",
    449: "Orientation and Position", 450: "BNO085 orientation and calibrated buoy motion", 451: "GPS location and UTC timing",
    452: "Bar02 pressure variation for wave-estimation research", 454: "Power and Thermal Sensors",
    455: "INA260 battery voltage and current", 456: "Second INA260 solar voltage and current",
    457: "MCP9808 electronics-enclosure temperature", 458: "Optional DS18B20 water temperature",
    460: "System-Health Inputs", 461: "Sensor validity, freshness, and sequence status", 462: "Internal enclosure temperature",
    463: "Battery and solar electrical status", 464: "GPS time and local system timestamps",
    477: "Runs on the Orange Pi Zero 3 target; a development laptop currently represents this role where hardware integration is incomplete.",
    501: "Orange Pi Zero 3 / Development Edge Host", 504: "Orange Pi Zero 3 / Development Edge Host",
    516: "The focused AI workflow receives validated, time-aligned measurements from the ESP32 data pipeline. Phase 1 limits AI use to short-horizon wave-height prediction and sea-condition classification.",
    517: "Primary AI Inputs", 518: "Calibrated BNO085 motion features", 519: "Bar02 pressure-derived wave features",
    520: "Wind speed", 521: "Wind direction", 522: "Recent observed wave-height history",
    523: "Timestamps and sampling quality", 524: "GPS position for record context",
    525: "Battery and solar status for data-quality context", 526: "MCP9808 enclosure temperature for system context",
    527: "Validated missing-data and stale-data flags", 528: "The model must not use uncalibrated raw acceleration as direct wave height.",
    529: "Inputs are accepted only after timestamp, range, and quality checks.", 530: "AI Outputs",
    531: "Predicted wave height for a 5–15 minute horizon", 532: "Sea-condition classification: Calm, Moderate, or Rough",
    533: "Prediction confidence or uncertainty indicator", 534: "Input-quality and insufficient-history status",
    535: "Explainable contributing-feature summary when implemented",
    536: "The AI does not claim tsunami, earthquake, tide, or general weather forecasting.",
    537: "The current dashboard trend predictor is a presentation baseline and is not yet a field-trained or safety-validated oceanographic model.",
    538: "Model accuracy, calibration, and uncertainty must be evaluated using synchronized field observations and an appropriate reference instrument before operational claims are made.",
}


def replace_text(paragraph, text):
    """Keep paragraph formatting and the first run's character formatting."""
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def clear_text(paragraph):
    for run in paragraph.runs:
        run.text = ""


def add_figure(paragraph, image_path, caption, width=6.25):
    """Embed a centered figure and caption in an existing blank paragraph."""
    image_path = Path(image_path)
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    clear_text(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    caption_run = paragraph.add_run(f"\n{caption}")
    caption_run.bold = True
    caption_run.italic = True
    caption_run.font.size = Pt(9)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    for index, text in REPLACEMENTS.items():
        replace_text(document.paragraphs[index], text)
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace("Mini PC", "Orange Pi Zero 3")

    # Replace the old text-only overall flow with one readable visual workflow.
    for index in range(319, 351):
        clear_text(document.paragraphs[index])
    add_figure(
        document.paragraphs[319],
        VISUALS / "system-architecture.png",
        "Figure 1. Current local-first workflow: sensing, ESP32 acquisition, Orange Pi edge processing, and dashboard presentation.",
    )

    # Actual current Fusion views are used instead of generic or legacy buoy art.
    add_figure(
        document.paragraphs[379],
        SCREENSHOTS / "Screenshot (681).png",
        "Figure 2. Actual Project FALCON-01 V2 Fusion assembly showing the single-body float, rounded keel, tower, opposed solar panels, and sensor deck.",
    )
    add_figure(
        document.paragraphs[432],
        SCREENSHOTS / "Screenshot (689).png",
        "Figure 3. Actual V2 Fusion electronics placement inside the elevated rectangular marine pod.",
    )
    add_figure(
        document.paragraphs[465],
        SCREENSHOTS / "Screenshot (693).png",
        "Figure 4. Actual sealed-pod thermal assembly: two internal recirculation fans, component heat sinks, airflow guides, vertical thermal bridge, and rear external finned heat sink.",
    )
    add_figure(
        document.paragraphs[514],
        VISUALS / "dashboard-overview.png",
        "Figure 5. Current Dashboard Next interface for live measurements, wave forecast, sea-condition class, power status, GPS, and system health.",
    )
    document.core_properties.title = "Project FALCON-01 — Current Revision 5 Thesis Documentation"
    document.save(OUTPUT)

    check = Document(OUTPUT)
    full_text = "\n".join(p.text for p in check.paragraphs).lower()
    forbidden = ("four stabilizer", "four auxiliary stabilizer", "water leak sensor", "abnormal wave detection", "high tide detection", "low tide detection")
    found = [term for term in forbidden if term in full_text]
    if found:
        raise RuntimeError(f"Outdated terms remain: {found}")
    print(f"Revised thesis: {OUTPUT}")
    print(f"Preserved: {len(check.paragraphs)} paragraphs, {len(check.tables)} table, {len(check.sections)} section")
    print(f"Embedded figures: {len(check.inline_shapes)}")


if __name__ == "__main__":
    main()
