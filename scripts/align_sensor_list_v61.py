"""LEGACY v6.1 updater retained only for traceability after Bay Station v7.0."""

from pathlib import Path

from docx import Document


DOCUMENT = Path(__file__).resolve().parents[1] / "THESIS DOCUMENTATION" / "Project_FALCON_Sensor_List.docx"


def set_row(row, values):
    for cell, value in zip(row.cells, values):
        cell.text = value


def align():
    document = Document(DOCUMENT)
    paragraphs = document.paragraphs
    paragraphs[0].text = "Project FALCON Sensor and Monitoring Baseline — v6.1"
    paragraphs[2].text = (
        "This document follows docs/PROJECT_CONTEXT.md v6.1. It separates the required Phase 1 baseline "
        "from supporting, health, security, and future functions. Physical integration and calibration remain pending."
    )
    paragraphs[4].text = "The required Phase 1 sensing functions are water pressure, GPS position, wind speed, and wind direction."
    paragraphs[5].text = (
        "A Holykell HPT604 Type A 0–2 mH2O 4–20 mA candidate is the primary wave-observation input. Bar02 is bench-only. The system retains raw and "
        "filtered pressure, baseline, optional depth, estimated wave height, validity, timestamp/source, and calibration state."
    )
    paragraphs[16].text = (
        "Required and supporting inputs → ESP32 acquisition, calibration, validation, and timestamping → USB serial/UART "
        "→ Orange Pi Zero 3 4GB → SQLite and deterministic local processing → REST API → local dashboard. "
        "Internet synchronization and AI are optional and non-blocking."
    )
    paragraphs[19].text = (
        "Pressure-derived wave height must be labeled ESTIMATED and CALIBRATION REQUIRED until controlled validation is complete."
    )
    paragraphs[22].text = "AI is optional and supporting; core monitoring, logging, security, and visualization must work without it."
    paragraphs[26].text = (
        "Phase 1 requires pressure, GPS, wind speed, and wind direction. A sealed DS18B20 provides supporting water "
        "temperature. Battery, solar, and enclosure temperature are health channels. Security uses persistent GPS "
        "geofence, a generic vibration/tamper input, an enclosure reed or limit switch, and a buzzer. BNO085, load-cell/HX711 "
        "mooring tension, salinity/conductivity, and mandatory AI are not required baseline functions."
    )

    main = document.tables[0]
    rows = [
        ["1", "Water-pressure sensor", "Holykell HPT604 Type A candidate; exact order code to be verified", "4–20 mA via protected ADS1115 receiver", "Raw/filtered pressure and calibrated estimated-wave support"],
        ["2", "GPS receiver", "Exact model TBD", "UART", "Position, UTC, fix validity, satellite state, and persistent geofence"],
        ["3", "Wind-speed sensor", "Exact model TBD", "Pulse/GPIO", "Local wind speed"],
        ["4", "Wind-direction sensor", "Exact model TBD", "Analog/ADC", "Local wind direction"],
    ]
    for row, values in zip(main.rows[1:], rows):
        set_row(row, values)
    last = main.rows[-1]._tr
    last.getparent().remove(last)

    support = document.tables[1]
    support_rows = [
        ["1", "Battery health monitor", "Exact design TBD", "Electrical interface TBD", "Battery voltage/current/power and charge state"],
        ["2", "Solar health monitor", "Exact design TBD", "Electrical interface TBD", "Solar charging and power status"],
        ["3", "Enclosure-temperature sensor", "Exact model TBD", "Interface TBD", "Thermal status and warning input"],
        ["4", "Water-temperature probe", "Sealed DS18B20", "One-Wire", "Supporting water-temperature context"],
    ]
    for row, values in zip(support.rows[1:], support_rows):
        set_row(row, values)

    future = document.tables[3]
    set_row(future.rows[1], ["BNO085 or other IMU", "Optional motion research only; not required for wave estimation", "Separate approved objective, calibration, and validation"])
    for row in future.rows[2:]:
        if "Leak or water-ingress" in row.cells[0].text or "Humidity/condensation" in row.cells[0].text:
            row.cells[2].text = "Not in current baseline; requires approved scope and controlled validation"

    controllers = document.tables[4]
    controllers.rows[2].cells[1].text = (
        "Receives ESP32 telemetry, stores SQLite records, performs deterministic local processing, and hosts the REST API and dashboard; optional AI remains isolated and non-blocking"
    )

    document.core_properties.subject = "Project FALCON-01 v6.1 sensor, health, and security baseline"
    document.core_properties.comments = "Aligned to docs/PROJECT_CONTEXT.md v6.1 on 2026-08-27."
    temporary = DOCUMENT.with_suffix(".aligned.docx")
    document.save(temporary)
    temporary.replace(DOCUMENT)
    print(DOCUMENT)


if __name__ == "__main__":
    raise SystemExit("Legacy updater blocked: align future records with PROJECT_CONTEXT.md v7.0")
